"""Снимок учёта из Core в XLSX (три вкладки) — для загрузки в Google Drive, когда прямой
записи в таблицу нет. Числа — числами (без «.00»), шапка и колонка «Имя» закреплены,
технические колонки скрыты. python build_core_sheet_xlsx.py <out.xlsx>"""
from __future__ import annotations
import re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from wwc_sql import run_sql

HIDDEN = {"chat_id", "owner_id", "payment_id", "entry_id"}
NUM = re.compile(r"^-?\d+(\.\d+)?$")

def cell(v):
    if v is None: return ""
    if isinstance(v, str) and NUM.match(v): return float(v) if "." in v else int(v)
    return v

def build(out: str) -> None:
    text = open("sheet_sync_queries.sql", encoding="utf-8").read()
    parts = re.split(r"^-- @tab (.+)$", text, flags=re.M)
    wb = Workbook(); wb.remove(wb.active)
    for i in range(1, len(parts), 2):
        name, sql = parts[i].strip(), parts[i + 1]
        rows = run_sql(sql); ws = wb.create_sheet(name)
        if not rows: continue
        cols = list(rows[0].keys()); ws.append(cols)
        for r in rows: ws.append([cell(r.get(c)) for c in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3B34"); c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "B2"
        for j, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(10, min(48, max(len(str(c)), *(len(str(r.get(c) or "")) for r in rows)) + 2))
            if c in HIDDEN: ws.column_dimensions[get_column_letter(j)].hidden = True
        ws.auto_filter.ref = ws.dimensions
        print(name, len(rows), "строк")
    wb.save(out); print("saved", out)

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    build(sys.argv[1] if len(sys.argv) > 1 else "WWC_uchet_core.xlsx")
